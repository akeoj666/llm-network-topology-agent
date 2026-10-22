"""把 pptx 的全部文字（含文本框、表格、组合内形状、备注页）抽成纯文本。

用法：
    python tools\\extract_pptx_text.py "C:\\path\\a.pptx" [-o out\\a.txt]

为什么单独写这个：课件里的关键信息（命令行、红字注释、步骤编号）大量放在
**独立文本框**里，不是占位符正文。只读 `slide.shapes.placeholders` 或
只读 title/body 会漏掉一半内容，看起来就像「读不到」。

所以这里做三件事：
1. 递归进 GroupShape（组合里的文本框也要拿到）
2. 读 Table 的每个单元格
3. 读 notes_slide 的备注（老师常把操作细节写在备注里）

输出格式按「幻灯片序号 + 形状类型」分块，方便回查某一句出自哪一页。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _shape_lines(shape, depth: int = 0) -> list[str]:
    """把一个形状（可能是组合）里的文字全部摊平。"""
    pad = "  " * depth
    out: list[str] = []

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            out.extend(_shape_lines(sub, depth + 1))
        return out

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.replace("\n", " ⏎ ").strip() for c in row.cells]
            if any(cells):
                out.append(f"{pad}[表格行] " + " | ".join(cells))
        return out

    if getattr(shape, "has_text_frame", False):
        for para in shape.text_frame.paragraphs:
            text = "".join(r.text for r in para.runs).strip()
            if not text:
                continue
            # 标出这是文本框（不是占位符），PPT 里红字注释基本都在这里
            kind = "文本框" if shape.is_placeholder is False else "占位符"
            out.append(f"{pad}[{kind}] {text}")
    return out


def extract(path: Path) -> str:
    prs = Presentation(str(path))
    lines: list[str] = [
        f"# {path.name}",
        f"# 共 {len(prs.slides)} 页，尺寸 {prs.slide_width}×{prs.slide_height} EMU",
        "",
    ]
    for idx, slide in enumerate(prs.slides, start=1):
        block: list[str] = []
        for shape in slide.shapes:
            block.extend(_shape_lines(shape))
        # 备注页
        try:
            note = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, KeyError):
            note = ""
        if note:
            block.append(f"[备注] {note.replace(chr(10), ' ⏎ ')}")
        if not block:
            continue
        lines.append(f"=== 第 {idx} 页 ===")
        lines.extend(block)
        lines.append("")
    return "\n".join(lines)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", default="")
    args = ap.parse_args()

    src = Path(args.pptx)
    if not src.exists():
        print(f"[错误] 找不到 {src}", file=sys.stderr)
        return 1

    text = extract(src)
    if args.out:
        dst = Path(args.out)
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_text(text, encoding="utf-8")
        print(f"已写入 {dst}（{len(text)} 字符）")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
